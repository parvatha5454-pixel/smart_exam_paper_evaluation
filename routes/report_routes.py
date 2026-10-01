"""
routes/report_routes.py
--------------------------
Module 6: Reports
- Per-student PDF mark sheet (one student, one exam, question-wise
  marks + total).
- Whole-class Excel summary (one exam, every evaluated student as a
  row, one column per question + total + percentage).

Both are generated on demand (not pre-built), saved into
Config.REPORTS_FOLDER, logged in the `reports` table for a history
trail, and sent to the browser as a download.
"""

import os
from datetime import datetime

from flask import Blueprint, render_template, request, redirect, url_for, flash, session, send_file

from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

from config import Config
from models.answer_key_model import get_all_answer_keys, get_answer_key_by_id, get_questions_for_answer_key
from models.scanned_sheet_model import get_all_scanned_sheets, get_scanned_sheet_by_id
from models.evaluation_model import get_evaluations_for_sheet, get_evaluation_count
from models.report_model import create_report
from utils.security import login_required
from utils.logger import log_activity

report_bp = Blueprint("report", __name__, url_prefix="/reports")


# ---------------------------------------------------------------
# Reports: pick an exam, see its evaluated students, download
# individual PDFs or the whole-class Excel summary
# ---------------------------------------------------------------

@report_bp.route("/")
@login_required
def list_reports():
    answer_keys = get_all_answer_keys()
    selected_id = request.args.get("answer_key_id", type=int)

    sheets = []
    key = None
    if selected_id:
        key = get_answer_key_by_id(selected_id)
        all_sheets = get_all_scanned_sheets(answer_key_id=selected_id)
        for s in all_sheets:
            if get_evaluation_count(s["sheet_id"]) == 0:
                continue
            results = get_evaluations_for_sheet(s["sheet_id"])
            s["total_awarded"] = sum(float(r["marks_awarded"] or 0) for r in results)
            s["total_max"] = sum(float(r["max_marks"] or 0) for r in results)
            s["percentage"] = (
                round((s["total_awarded"] / s["total_max"]) * 100, 1) if s["total_max"] else 0
            )
            sheets.append(s)

    return render_template(
        "reports/list.html",
        answer_keys=answer_keys,
        selected_id=selected_id,
        key=key,
        sheets=sheets,
    )


# ---------------------------------------------------------------
# Individual student PDF mark sheet
# ---------------------------------------------------------------

@report_bp.route("/<int:sheet_id>/pdf")
@login_required
def generate_pdf_report(sheet_id):
    sheet = get_scanned_sheet_by_id(sheet_id)
    if not sheet:
        flash("Scanned sheet not found.", "danger")
        return redirect(url_for("report.list_reports"))

    results = get_evaluations_for_sheet(sheet_id)
    if not results:
        flash("This sheet hasn't been evaluated yet - nothing to report.", "warning")
        return redirect(url_for("report.list_reports", answer_key_id=sheet["answer_key_id"]))

    total_awarded = sum(float(r["marks_awarded"] or 0) for r in results)
    total_max = sum(float(r["max_marks"] or 0) for r in results)
    percentage = round((total_awarded / total_max) * 100, 1) if total_max else 0

    os.makedirs(Config.REPORTS_FOLDER, exist_ok=True)
    filename = f"marksheet_{sheet['roll_number']}_{sheet['answer_key_id']}_{sheet_id}.pdf"
    file_path = os.path.join(Config.REPORTS_FOLDER, filename)

    _build_pdf_marksheet(file_path, sheet, results, total_awarded, total_max, percentage)

    create_report(sheet_id, "pdf", file_path, total_awarded, percentage)
    log_activity(
        session.get("admin_id"),
        "GENERATE_REPORT",
        f"Generated PDF mark sheet for sheet #{sheet_id} ({sheet['student_name']})",
    )

    return send_file(file_path, as_attachment=True, download_name=filename)


def _build_pdf_marksheet(file_path, sheet, results, total_awarded, total_max, percentage):
    doc = SimpleDocTemplate(file_path, pagesize=A4, topMargin=2 * cm, bottomMargin=2 * cm)
    styles = getSampleStyleSheet()
    elements = []

    elements.append(Paragraph("Exam Answer Analysis - Mark Sheet", styles["Title"]))
    elements.append(Spacer(1, 0.5 * cm))

    info_data = [
        ["Student Name:", sheet["student_name"], "Roll No:", sheet["roll_number"]],
        ["Exam:", sheet["exam_name"], "Subject:", f"{sheet['subject_name']}"],
        ["Generated:", datetime.now().strftime("%Y-%m-%d %H:%M"), "", ""],
    ]
    info_table = Table(info_data, colWidths=[3 * cm, 6 * cm, 3 * cm, 5 * cm])
    info_table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
                ("FONTNAME", (2, 0), (2, -1), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 10),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    elements.append(info_table)
    elements.append(Spacer(1, 0.7 * cm))

    table_data = [["Q. No.", "Max Marks", "Marks Awarded", "Remarks"]]
    for r in results:
        table_data.append(
            [
                str(r["question_number"]),
                str(r["max_marks"]),
                str(r["marks_awarded"]),
                r["remarks"] or "",
            ]
        )
    table_data.append(["", "Total", f"{total_awarded} / {total_max}", f"{percentage}%"])

    marks_table = Table(table_data, colWidths=[2.5 * cm, 3 * cm, 3.5 * cm, 7 * cm])
    marks_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f2937")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
                ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#e5e7eb")),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("ALIGN", (1, 0), (2, -1), "CENTER"),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    elements.append(marks_table)

    doc.build(elements)


# ---------------------------------------------------------------
# Whole-class Excel summary for one exam
# ---------------------------------------------------------------

@report_bp.route("/summary/<int:answer_key_id>/excel")
@login_required
def generate_summary_excel(answer_key_id):
    key = get_answer_key_by_id(answer_key_id)
    if not key:
        flash("Answer key not found.", "danger")
        return redirect(url_for("report.list_reports"))

    questions = get_questions_for_answer_key(answer_key_id)
    all_sheets = get_all_scanned_sheets(answer_key_id=answer_key_id)
    evaluated_sheets = [s for s in all_sheets if get_evaluation_count(s["sheet_id"]) > 0]

    if not evaluated_sheets:
        flash("No evaluated sheets for this exam yet - nothing to summarize.", "warning")
        return redirect(url_for("report.list_reports", answer_key_id=answer_key_id))

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"

    header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    headers = ["Roll No.", "Student Name"] + [f"Q{q['question_number']}" for q in questions] + [
        "Total",
        "Max",
        "Percentage",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    total_max = sum(float(q["max_marks"] or 0) for q in questions)

    for s in evaluated_sheets:
        results = get_evaluations_for_sheet(s["sheet_id"])
        marks_by_question = {r["question_id"]: r["marks_awarded"] for r in results}
        row = [s["roll_number"], s["student_name"]]
        for q in questions:
            row.append(float(marks_by_question.get(q["question_id"], 0) or 0))
        total_awarded = sum(row[2:])
        percentage = round((total_awarded / total_max) * 100, 1) if total_max else 0
        row += [total_awarded, total_max, f"{percentage}%"]
        ws.append(row)

    for col in ws.columns:
        max_len = max(len(str(c.value)) for c in col if c.value is not None)
        ws.column_dimensions[col[0].column_letter].width = max(10, max_len + 2)

    os.makedirs(Config.REPORTS_FOLDER, exist_ok=True)
    safe_exam_name = "".join(c if c.isalnum() else "_" for c in key["exam_name"])
    filename = f"summary_{safe_exam_name}_{answer_key_id}.xlsx"
    file_path = os.path.join(Config.REPORTS_FOLDER, filename)
    wb.save(file_path)

    create_report(evaluated_sheets[0]["sheet_id"], "excel", file_path, None, None)
    log_activity(
        session.get("admin_id"),
        "GENERATE_REPORT",
        f"Generated class summary Excel for answer key #{answer_key_id} "
        f"({len(evaluated_sheets)} student(s))",
    )

    return send_file(file_path, as_attachment=True, download_name=filename)
